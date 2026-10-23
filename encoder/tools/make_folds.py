"""Generate encoder fold assignments from an APR workbook.

The random mode reproduces the released row-level K-fold protocol. The group
mode accepts a separate CSV mapping ``sequence_id,group_id``; this keeps
homology clustering and its large intermediate files outside the repository.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import GroupKFold, KFold


def read_sequences(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook["sequences"].iter_rows(values_only=True)
        header = [str(value) for value in next(rows)]
        records = [dict(zip(header, row)) for row in rows]
    finally:
        workbook.close()
    if not records:
        raise ValueError("The sequences sheet is empty")
    ids = [str(row["sequence_id"]) for row in records]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate sequence_id values")
    return records


def read_groups(path: Path, ids: set[str]) -> dict[str, str]:
    with path.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    result = {row["sequence_id"]: row["group_id"] for row in rows}
    if set(result) != ids or any(not value for value in result.values()):
        raise ValueError("group CSV must contain exactly sequence_id,group_id for every sequence")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate APR encoder fold assignments")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--folds", type=int, default=10)
    parser.add_argument("--seed", type=int, default=2026)
    parser.add_argument("--group-csv", type=Path,
                        help="Optional sequence_id,group_id mapping for grouped folds")
    args = parser.parse_args()
    if args.folds < 2:
        raise ValueError("folds must be at least 2")
    records = read_sequences(args.workbook)
    ids = [str(row["sequence_id"]) for row in records]
    if args.group_csv:
        groups = read_groups(args.group_csv, set(ids))
        splitter = GroupKFold(n_splits=args.folds)
        iterator = splitter.split(ids, groups=[groups[item] for item in ids])
        method = "sklearn.model_selection.GroupKFold"
    else:
        splitter = KFold(n_splits=args.folds, shuffle=True, random_state=args.seed)
        iterator = splitter.split(np.arange(len(ids)))
        method = "sklearn.model_selection.KFold"
    assignments = np.full(len(ids), -1, dtype=np.int64)
    for fold, (_, held_out) in enumerate(iterator):
        assignments[held_out] = fold
    if np.any(assignments < 0):
        raise RuntimeError("Some records were not assigned to a fold")

    args.output_root.mkdir(parents=True, exist_ok=True)
    with (args.output_root / "apr_folds.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["dataset", "item_id", "fold", "label", "sample_class"])
        for row, fold in zip(records, assignments.tolist()):
            writer.writerow(["apr", str(row["sequence_id"]), int(fold),
                             row.get("positive_residue_count", ""), "apr_region_sequence"])
    manifest = {
        "folds": args.folds, "seed": args.seed, "method": method,
        "shuffle": args.group_csv is None, "indexing": "zero_based_in_manifest",
        "datasets": {"apr": {"count": len(records), "id_column": "sequence_id",
                              "group_column": "group_id" if args.group_csv else None,
                              "fold_by_id": {item: int(fold) for item, fold in zip(ids, assignments.tolist())}}}
    }
    (args.output_root / "fold_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"output_root": str(args.output_root), "records": len(records),
                      "folds": args.folds, "method": method}, indent=2))


if __name__ == "__main__":
    main()
