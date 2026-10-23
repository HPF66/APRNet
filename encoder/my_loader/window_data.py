"""APR sequence records, chunking, cached inputs and residue-aligned collation."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
import random
from dataclasses import dataclass
from pathlib import Path
import numpy as np
import torch
from openpyxl import load_workbook
from sklearn.metrics import average_precision_score, matthews_corrcoef, roc_auc_score
from torch import nn
from torch.nn import functional as F
from torch.utils.data import DataLoader, Dataset
WINDOW_WIDTH = 320
ESM_WIDTH = 480
WINDOW_SCALES = (6, 8, 10, 12, 16, 20)
EVIDENCE_FEATURES = len(WINDOW_SCALES) * 4

@dataclass(frozen=True)
class Protein:
    sample_id: str
    sequence: str
    target: np.ndarray
    fold: int

@dataclass(frozen=True)
class Chunk:
    sample_id: str
    start: int
    stop: int
    fold: int

def require(ok: bool, message: str) -> None:
    if not ok:
        raise ValueError(message)

def set_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    torch.backends.cudnn.benchmark = False
    torch.backends.cudnn.deterministic = True

def read_fold_map(path: Path) -> dict[str, int]:
    with path.open(newline='', encoding='utf-8') as handle:
        rows = list(csv.DictReader(handle))
    result = {row['item_id']: int(row['fold']) for row in rows if row.get('dataset') == 'apr'}
    require(result, f'No APR folds in {path}')
    return result

def read_proteins(workbook_path: Path, fold_csv: Path) -> list[Protein]:
    fold_map = read_fold_map(fold_csv)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = workbook['sequences'].iter_rows(values_only=True)
    header = [str(v) for v in next(rows)]
    proteins: list[Protein] = []
    for row in rows:
        item = dict(zip(header, row))
        sample_id = str(item['sequence_id'])
        sequence = str(item.get('sequence') or '').strip().upper()
        if not sequence:
            continue
        require(sample_id in fold_map, f'Missing APR fold for {sample_id}')
        mask = str(item.get('apr_positive_mask') or '').strip()
        require(len(mask) == len(sequence), f'APR mask length mismatch for {sample_id}')
        require(not set(mask) - {'0', '1', '?'}, f'Unsupported APR mask for {sample_id}')
        target = np.fromiter((1 if char == '1' else 0 for char in mask), dtype=np.int8, count=len(sequence))
        proteins.append(Protein(sample_id, sequence, target, fold_map[sample_id]))
    workbook.close()
    require(proteins, 'No APR sequences')
    require(len({p.sample_id for p in proteins}) == len(proteins), 'Duplicate APR IDs')
    return proteins


def _index_manifest(root: Path) -> dict:
    manifest = json.loads((root / 'manifest.json').read_text(encoding='utf-8'))
    indexed = {}
    for record in manifest['records']:
        sid = record['sample_id']
        require(sid not in indexed, f'Duplicate cache manifest ID: {sid} in {root}')
        indexed[sid] = record
    return indexed


def _cache_path(root: Path, record: dict, protein: Protein) -> Path:
    require(record.get('sequence') == protein.sequence, f'Cache sequence mismatch: {protein.sample_id} in {root}')
    require(record.get('length') == len(protein.sequence), f'Cache length metadata mismatch: {protein.sample_id}')
    filename = record['cache_file']
    require(Path(filename).name == filename and chr(92) not in filename,
            f'Invalid cache filename: {filename}')
    return root / 'samples' / filename


def _npz_identity(z, protein: Protein, required=False) -> None:
    for key, expected in (('sample_id', protein.sample_id), ('sequence', protein.sequence)):
        require(not required or key in z.files, f'Cache lacks {key}: {protein.sample_id}')
        if key in z.files:
            require(z[key].ndim == 0 and str(z[key].item()) == expected, f'NPZ {key} mismatch: {protein.sample_id}')


def _integer(z, key: str) -> np.ndarray:
    value = z[key]
    require(value.ndim == 1 and np.issubdtype(value.dtype, np.integer), f'{key} must be a vector of integers')
    return value.astype(np.int64)


def load_store(proteins: list[Protein], window_root: Path, esm_root: Path) -> dict[str, dict[str, np.ndarray]]:
    window_by_id = _index_manifest(window_root)
    esm_by_id = _index_manifest(esm_root)
    require(len({p.sample_id for p in proteins}) == len(proteins), 'Duplicate requested protein IDs')
    store: dict[str, dict[str, np.ndarray]] = {}
    for protein in proteins:
        sid, n = protein.sample_id, len(protein.sequence)
        require(n > 0, f'Empty sequence: {sid}')
        require(sid in window_by_id, f'Missing window cache: {sid}')
        require(sid in esm_by_id, f'Missing ESM cache: {sid}')
        with np.load(_cache_path(window_root, window_by_id[sid], protein), allow_pickle=False) as z:
            _npz_identity(z, protein, required=True)
            starts = _integer(z, 'window_start')
            lengths = _integer(z, 'window_length')
            embeddings = z['window_embedding'].astype(np.float32)
            probabilities = z['window_probability'].astype(np.float32)
            local_probabilities = z['window_local_probability'].astype(np.float32) if 'window_local_probability' in z.files else probabilities.copy()
            scale_ids = _integer(z, 'window_scale_id') if 'window_scale_id' in z.files else np.asarray([WINDOW_SCALES.index(int(x)) if int(x) in WINDOW_SCALES else 0 if n < WINDOW_SCALES[0] and int(x) == n else -1 for x in lengths], dtype=np.int64)
            count = len(starts)
            require(count > 0 and lengths.shape == probabilities.shape == local_probabilities.shape == scale_ids.shape == (count,), f'Window array mismatch: {sid}')
            require(embeddings.shape == (count, WINDOW_WIDTH), f'Window width/count mismatch: {sid}')
            require(np.all((starts >= 0) & (lengths > 0) & (starts + lengths <= n)), f'Window bounds mismatch: {sid}')
            require(np.all((scale_ids >= 0) & (scale_ids < len(WINDOW_SCALES))), f'Unsupported window scales: {sid}')
            require(np.all(lengths == np.minimum(np.asarray(WINDOW_SCALES)[scale_ids], n)), f'Window length/scale mismatch: {sid}')
            require(len(set(zip(starts.tolist(), lengths.tolist(), scale_ids.tolist()))) == count, f'Duplicate windows: {sid}')
            for value in (embeddings, probabilities, local_probabilities):
                require(np.isfinite(value).all(), f'Nonfinite window array: {sid}')
            for value in (probabilities, local_probabilities):
                require(np.all((value >= 0) & (value <= 1)), f'Probability outside [0,1]: {sid}')
            if 'window_count' in window_by_id[sid]:
                require(window_by_id[sid]['window_count'] == count, f'Window count metadata mismatch: {sid}')
            profile_hidden = profile_local = None
            profile_keys = {'profile_offsets', 'profile_hidden', 'profile_local_probability'}
            present = profile_keys.intersection(z.files)
            require(not present or present == profile_keys, f'Incomplete profile cache: {sid}')
            if present:
                offsets = _integer(z, 'profile_offsets')
                flat_hidden = z['profile_hidden'].astype(np.float32)
                flat_local = z['profile_local_probability'].astype(np.float32)
                require(offsets.shape == (count + 1,) and offsets[0] == 0 and np.array_equal(np.diff(offsets), lengths), f'Profile offset mismatch: {sid}')
                require(flat_hidden.shape == (offsets[-1], WINDOW_WIDTH) and flat_local.shape == (offsets[-1],), f'Profile shape mismatch: {sid}')
                require(np.isfinite(flat_hidden).all() and np.isfinite(flat_local).all() and np.all((flat_local >= 0) & (flat_local <= 1)), f'Invalid profile values: {sid}')
                profile_hidden = np.zeros((n, WINDOW_WIDTH), dtype=np.float32)
                profile_local = np.zeros(n, dtype=np.float32)
                profile_counts = np.zeros(n, dtype=np.float32)
                for wi, (start, length) in enumerate(zip(starts, lengths)):
                    begin, end = offsets[wi:wi + 2]
                    profile_hidden[start:start + length] += flat_hidden[begin:end]
                    profile_local[start:start + length] += flat_local[begin:end]
                    profile_counts[start:start + length] += 1.0
                profile_hidden /= np.maximum(profile_counts[:, None], 1.0)
                profile_local /= np.maximum(profile_counts, 1.0)
        with np.load(_cache_path(esm_root, esm_by_id[sid], protein), allow_pickle=False) as z:
            _npz_identity(z, protein)
            esm = z['hidden'].astype(np.float32)
        require(esm.shape == (n, ESM_WIDTH) and np.isfinite(esm).all(), f'Invalid ESM shape/values: {sid}')
        covered = np.zeros(n, dtype=np.int32)
        for start, length in zip(starts, lengths):
            covered[start:start + length] += 1
        require(np.all(covered > 0), f'Window cache does not cover sequence: {sid}')
        item = {'esm': esm, 'starts': starts.astype(np.int32), 'lengths': lengths.astype(np.int32), 'tokens': embeddings, 'window_probability': probabilities, 'window_local_probability': local_probabilities, 'window_scale_id': scale_ids.astype(np.int32)}
        if profile_hidden is not None:
            item['profile_hidden'] = profile_hidden
            item['profile_local_probability'] = profile_local
        store[sid] = item
    return store

def make_chunks(proteins: list[Protein], max_length: int, overlap: int) -> list[Chunk]:
    require(max_length > 0 and 0 <= overlap < max_length, 'invalid chunk geometry')
    step = max_length - overlap
    chunks: list[Chunk] = []
    for protein in proteins:
        starts = list(range(0, max(len(protein.sequence) - max_length, 0) + 1, step))
        if not starts or starts[-1] + max_length < len(protein.sequence):
            starts.append(max(0, len(protein.sequence) - max_length))
        for start in sorted(set(starts)):
            chunks.append(Chunk(protein.sample_id, start, min(start + max_length, len(protein.sequence)), protein.fold))
    return chunks

class ChunkDataset(Dataset[Chunk]):

    def __init__(self, chunks: list[Chunk], proteins: dict[str, Protein], store: dict[str, dict[str, np.ndarray]]):
        self.chunks, self.proteins, self.store = (chunks, proteins, store)

    def __len__(self) -> int:
        return len(self.chunks)

    def __getitem__(self, index: int) -> Chunk:
        return self.chunks[index]

def build_residue_evidence(starts: np.ndarray, lengths: np.ndarray, probabilities: np.ndarray, local_probabilities: np.ndarray, scale_ids: np.ndarray, sequence_length: int) -> np.ndarray:
    """Build overlap-normalized multi-scale window evidence for one chunk.

    The soft-OR uses the geometric mean of (1 - p) across overlapping windows;
    this prevents dense overlapping windows from saturating the evidence merely
    because the same residue was covered many times.
    """
    evidence = np.zeros((sequence_length, EVIDENCE_FEATURES), dtype=np.float32)
    for scale_index in range(len(WINDOW_SCALES)):
        scale_mask = scale_ids == scale_index
        if not np.any(scale_mask):
            continue
        sum_log_not_p = np.zeros(sequence_length, dtype=np.float64)
        sum_log_not_local = np.zeros(sequence_length, dtype=np.float64)
        counts = np.zeros(sequence_length, dtype=np.int32)
        max_p = np.zeros(sequence_length, dtype=np.float32)
        max_local = np.zeros(sequence_length, dtype=np.float32)
        for start, length, probability, local_probability in zip(starts[scale_mask], lengths[scale_mask], probabilities[scale_mask], local_probabilities[scale_mask]):
            left = max(int(start), 0)
            right = min(sequence_length, int(start + length))
            if right <= left:
                continue
            sl = slice(left, right)
            probability = float(np.clip(probability, 0.0, 1.0 - 1e-06))
            local_probability = float(np.clip(local_probability, 0.0, 1.0 - 1e-06))
            sum_log_not_p[sl] += np.log1p(-probability)
            sum_log_not_local[sl] += np.log1p(-local_probability)
            counts[sl] += 1
            max_p[sl] = np.maximum(max_p[sl], probability)
            max_local[sl] = np.maximum(max_local[sl], local_probability)
        valid = counts > 0
        soft_or = np.zeros(sequence_length, dtype=np.float32)
        soft_or_local = np.zeros(sequence_length, dtype=np.float32)
        soft_or[valid] = 1.0 - np.exp(sum_log_not_p[valid] / counts[valid])
        soft_or_local[valid] = 1.0 - np.exp(sum_log_not_local[valid] / counts[valid])
        offset = scale_index * 4
        evidence[:, offset] = soft_or
        evidence[:, offset + 1] = soft_or_local
        evidence[:, offset + 2] = max_p
        evidence[:, offset + 3] = max_local
    return evidence

def collate(rows: list[Chunk], dataset: ChunkDataset) -> dict[str, object]:
    lengths = [r.stop - r.start for r in rows]
    tmax = max(lengths)
    selected: list[tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]] = []
    for row in rows:
        s = dataset.store[row.sample_id]
        valid = (s['starts'] < row.stop) & (s['starts'] + s['lengths'] > row.start)
        selected.append((s['tokens'][valid], s['starts'][valid] - row.start, s['lengths'][valid], s['window_probability'][valid], s['window_local_probability'][valid], s['window_scale_id'][valid]))
    wmax = max((item[0].shape[0] for item in selected))
    base = torch.zeros(len(rows), tmax, ESM_WIDTH)
    tokens = torch.zeros(len(rows), wmax, WINDOW_WIDTH)
    starts = torch.zeros(len(rows), wmax, dtype=torch.float32)
    token_lengths = torch.zeros(len(rows), wmax, dtype=torch.float32)
    token_prob = torch.zeros(len(rows), wmax)
    token_local_prob = torch.zeros(len(rows), wmax)
    token_scale = torch.zeros(len(rows), wmax)
    token_mask = torch.zeros(len(rows), wmax, dtype=torch.bool)
    residue_evidence = torch.zeros(len(rows), tmax, EVIDENCE_FEATURES)
    profile_hidden = torch.zeros(len(rows), tmax, WINDOW_WIDTH)
    profile_local = torch.zeros(len(rows), tmax)
    target = torch.zeros(len(rows), tmax)
    residue_mask = torch.zeros(len(rows), tmax, dtype=torch.bool)
    ids: list[str] = []
    chunk_starts: list[int] = []
    for i, row in enumerate(rows):
        n = row.stop - row.start
        protein = dataset.proteins[row.sample_id]
        s = dataset.store[row.sample_id]
        base[i, :n] = torch.from_numpy(s['esm'][row.start:row.stop])
        target[i, :n] = torch.from_numpy(protein.target[row.start:row.stop].astype(np.float32))
        residue_mask[i, :n] = True
        tok, st, ln, pr, local_pr, scale_ids = selected[i]
        m = len(tok)
        tokens[i, :m] = torch.from_numpy(tok)
        starts[i, :m] = torch.from_numpy(st.astype(np.float32))
        token_lengths[i, :m] = torch.from_numpy(ln.astype(np.float32))
        token_prob[i, :m] = torch.from_numpy(pr)
        token_local_prob[i, :m] = torch.from_numpy(local_pr)
        token_scale[i, :m] = torch.from_numpy(scale_ids.astype(np.float32))
        token_mask[i, :m] = True
        residue_evidence[i, :n] = torch.from_numpy(build_residue_evidence(st, ln, pr, local_pr, scale_ids, n))
        if 'profile_hidden' in s:
            profile_hidden[i, :n] = torch.from_numpy(s['profile_hidden'][row.start:row.stop])
            profile_local[i, :n] = torch.from_numpy(s['profile_local_probability'][row.start:row.stop])
        ids.append(row.sample_id)
        chunk_starts.append(row.start)
    return {'base': base, 'tokens': tokens, 'token_starts': starts, 'token_lengths': token_lengths, 'token_probability': token_prob, 'token_local_probability': token_local_prob, 'token_scale': token_scale, 'token_mask': token_mask, 'residue_evidence': residue_evidence, 'profile_hidden': profile_hidden, 'profile_local_probability': profile_local, 'target': target, 'mask': residue_mask, 'sample_ids': ids, 'starts': chunk_starts}

def make_loader(chunks: list[Chunk], proteins: dict[str, Protein], store: dict[str, dict[str, np.ndarray]], batch_size: int, shuffle: bool, seed: int, workers: int) -> DataLoader:
    dataset = ChunkDataset(chunks, proteins, store)
    generator = torch.Generator().manual_seed(seed)
    return DataLoader(dataset, batch_size=batch_size, shuffle=shuffle, generator=generator, num_workers=workers, collate_fn=lambda rows: collate(rows, dataset))

def move(batch: dict[str, object], device: torch.device) -> dict[str, object]:
    return {k: v.to(device) if isinstance(v, torch.Tensor) else v for k, v in batch.items()}

def predict(model: nn.Module, chunks: list[Chunk], proteins: dict[str, Protein], store: dict[str, dict[str, np.ndarray]], args: argparse.Namespace, device: torch.device) -> dict[str, np.ndarray]:
    loader = make_loader(chunks, proteins, store, args.batch_size, False, args.seed, args.num_workers)
    model.eval()
    sums: dict[str, np.ndarray] = {}
    counts: dict[str, np.ndarray] = {}
    with torch.inference_mode():
        for raw in loader:
            batch = move(raw, device)
            values = torch.sigmoid(model(batch['base'], batch['tokens'], batch['token_starts'], batch['token_lengths'], batch['token_mask'], batch['mask'], batch['token_probability'], batch['token_local_probability'], batch['token_scale'], batch['residue_evidence'], profile_hidden=batch['profile_hidden'], profile_local_probability=batch['profile_local_probability'])).cpu().numpy()
            mask = batch['mask'].cpu().numpy()
            for row, sample_id in enumerate(batch['sample_ids']):
                start = int(batch['starts'][row])
                length = int(mask[row].sum())
                n = len(proteins[sample_id].sequence)
                if sample_id not in sums:
                    sums[sample_id] = np.zeros(n, dtype=np.float64)
                    counts[sample_id] = np.zeros(n, dtype=np.float64)
                sums[sample_id][start:start + length] += values[row, :length]
                counts[sample_id][start:start + length] += 1.0
    return {k: sums[k] / np.maximum(counts[k], 1.0) for k in sums}

def metrics(proteins: list[Protein], probabilities: dict[str, np.ndarray], threshold: float) -> dict[str, float | int]:
    y, p = ([], [])
    for protein in proteins:
        y.extend(protein.target.tolist())
        p.extend(probabilities[protein.sample_id].tolist())
    yv = np.asarray(y, dtype=np.int64)
    pv = np.asarray(p, dtype=np.float64)
    pred = pv >= threshold
    tp = int(np.sum((yv == 1) & pred))
    fn = int(np.sum((yv == 1) & ~pred))
    tn = int(np.sum((yv == 0) & ~pred))
    fp = int(np.sum((yv == 0) & pred))
    sens = tp / max(tp + fn, 1)
    spec = tn / max(tn + fp, 1)
    return {'n_residues': int(len(yv)), 'positive_residues': int(yv.sum()), 'negative_residues': int((yv == 0).sum()), 'mcc': float(matthews_corrcoef(yv, pred)), 'auroc': float(roc_auc_score(yv, pv)), 'auprc': float(average_precision_score(yv, pv)), 'sensitivity': float(sens), 'specificity': float(spec), 'balanced_accuracy': float((sens + spec) / 2), 'predicted_positive_rate': float(pred.mean()), 'true_positive': tp, 'false_negative': fn, 'true_negative': tn, 'false_positive': fp}

def calculate_pos_weight(chunks: list[Chunk], proteins: dict[str, Protein], device: torch.device, power: float) -> torch.Tensor:
    positive = sum((int(np.sum(proteins[c.sample_id].target[c.start:c.stop] == 1)) for c in chunks))
    negative = sum((int(np.sum(proteins[c.sample_id].target[c.start:c.stop] == 0)) for c in chunks))
    ratio = max(negative / max(positive, 1), 0.05)
    return torch.tensor(ratio ** power, dtype=torch.float32, device=device)
