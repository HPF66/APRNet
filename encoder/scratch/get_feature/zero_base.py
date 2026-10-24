#!/usr/bin/env python3
"""Create schema-compatible zero ESM cache for a deliberately ESM-free APR encoder.

``strict_c2_robust_v1.py`` validates the legacy ESM cache schema before
explicitly replacing every ESM tensor with zeros.  This script supplies that
placeholder schema for new proteins while recording that it carries no model
signal.  It must not be used by an ESM-consuming model.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ESM_WIDTH = 480


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apr-workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def cache_name(sample_id: str) -> str:
    return hashlib.sha256(sample_id.encode("utf-8")).hexdigest() + ".npz"


def main() -> None:
    args = parse_args()
    workbook = load_workbook(args.apr_workbook, read_only=True, data_only=True)
    rows = workbook["sequences"].iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]
    records = []
    for row in rows:
        item = dict(zip(header, row))
        sample_id = str(item["sequence_id"])
        sequence = str(item.get("sequence") or "").strip().upper()
        if sequence:
            records.append({"sample_id": sample_id, "sequence": sequence, "length": len(sequence)})
    workbook.close()
    if not records:
        raise ValueError("No sequences")
    args.output_dir.mkdir(parents=True, exist_ok=False)
    sample_dir = args.output_dir / "samples"
    sample_dir.mkdir()
    manifest = []
    for index, record in enumerate(records, start=1):
        filename = cache_name(record["sample_id"])
        np.savez_compressed(sample_dir / filename, hidden=np.zeros((record["length"], ESM_WIDTH), dtype=np.float16))
        manifest.append({**record, "cache_file": filename})
        print(json.dumps({"status": "zero_esm_cache", "completed": index, "total": len(records), "sample_id": record["sample_id"]}), flush=True)
    metadata = {
        "definition": "Zero placeholder only for strict ESM-free APR encoders",
        "encoder_signal": "none",
        "hidden_width": ESM_WIDTH,
        "sample_count": len(manifest),
        "records": manifest,
    }
    (args.output_dir / "manifest.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "complete", "output_dir": str(args.output_dir)}), flush=True)


if __name__ == "__main__":
    main()
