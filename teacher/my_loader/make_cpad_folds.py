"""Create the reproducible random K-fold manifest used by the CPAD teacher.

Only the source workbook is required. Fold outputs are intentionally kept
outside the repository because they are generated artifacts.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import KFold


def read_peptides(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["peptides"]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value) for value in next(rows)]
        records = [dict(zip(header, row)) for row in rows]
    finally:
        workbook.close()
    if not records:
        raise ValueError("The peptides sheet is empty")
    ids = [str(row["peptide_id"]) for row in records]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate peptide_id values")
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate CPAD peptide K-fold assignments")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--folds", type=int, default=10)
    parser.add_argument("--seed", type=int, default=2026)
    args = parser.parse_args()
    if args.folds < 2:
        raise ValueError("folds must be at least 2")
    records = read_peptides(args.workbook)
    assignments = np.full(len(records), -1, dtype=np.int64)
    splitter = KFold(n_splits=args.folds, shuffle=True, random_state=args.seed)
    for fold, (_, held_out) in enumerate(splitter.split(np.arange(len(records)))):
        assignments[held_out] = fold
    if np.any(assignments < 0):
        raise RuntimeError("Some records were not assigned to a fold")

    args.output_root.mkdir(parents=True, exist_ok=True)
    with (args.output_root / "peptide_folds.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["dataset", "item_id", "fold", "label"])
        for row, fold in zip(records, assignments.tolist()):
            writer.writerow(["peptide", str(row["peptide_id"]), int(fold), row.get("label", "")])

    manifest = {
        "folds": args.folds,
        "seed": args.seed,
        "method": "sklearn.model_selection.KFold",
        "shuffle": True,
        "indexing": "zero_based_in_manifest",
        "warning": "Row-level random baseline; related proteins may cross folds.",
        "datasets": {
            "peptide": {
                "count": len(records),
                "id_column": "peptide_id",
                "fold_by_id": {
                    str(row["peptide_id"]): int(fold)
                    for row, fold in zip(records, assignments.tolist())
                },
            }
        },
    }
    (args.output_root / "fold_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"output_root": str(args.output_root), "records": len(records),
                      "folds": args.folds, "seed": args.seed}, indent=2))


if __name__ == "__main__":
    main()
