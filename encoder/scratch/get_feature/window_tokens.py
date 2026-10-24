#!/usr/bin/env python3
"""Build independent Peptide window-token caches for APR transfer.

Unlike the legacy teacher cache, this file intentionally keeps one token per
sliding window.  No overlap-weighted residue averaging is performed here.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook

from encoder.utils.window_policy import enumerate_windows
from encoder.utils.lineage import MASKED_LOCAL_MEAN


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--backbone", type=Path, required=True)
    parser.add_argument("--peptide-checkpoint", type=Path, default=Path(__file__).resolve().parents[3] / "teacher/pretrained/teacher_canya_cpad.pt")
    parser.add_argument("--apr-workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--device", default="cuda")
    parser.add_argument("--window-lengths", default="6,8,10,12,16,20")
    parser.add_argument("--stride", type=int, default=1)
    parser.add_argument("--batch-size", type=int, default=128)
    return parser.parse_args()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def load_peptide_teacher(backbone: Path, checkpoint_path: Path, device: torch.device):
    """Restore the exact full-data Peptide Teacher checkpoint.

    The previous cache builder depended on an experiment-local extractor that
    is no longer present on the server.  The transfer checkpoint already
    contains its complete ``cross_validate`` argument set and model weights,
    so rebuilding it through the maintained project entry point is both more
    reproducible and independent of the removed experiment directory.
    """
    import teacher.train.cross_validate as cross_validate

    checkpoint = torch.load(checkpoint_path, map_location="cpu", weights_only=False)
    require(isinstance(checkpoint, dict) and "model" in checkpoint and "args" in checkpoint,
            f"Unsupported Peptide checkpoint format: {checkpoint_path}")
    saved_args = checkpoint["args"]
    if isinstance(saved_args, dict):
        args = argparse.Namespace(**saved_args)
    else:
        args = saved_args
    from teacher.utils.runtime import build_inference_model
    model = build_inference_model(args, device, backbone)
    incompatible = model.load_state_dict(checkpoint["model"], strict=False)
    # Transformers versions differ on whether an unused absolute-position
    # tensor is materialized for the rotary ESM2 backbone.  It is never read
    # by this checkpoint's rotary attention path, so tolerate that one legacy
    # source-only key while rejecting all real architecture mismatches.
    allowed_unexpected = {
        "sequence.encoder.backbone.backbone.base_model.model.embeddings.position_embeddings.weight"
    }
    unexpected = sorted(set(incompatible.unexpected_keys) - allowed_unexpected)
    require(not incompatible.missing_keys and not unexpected,
            "Peptide checkpoint/model mismatch: "
            f"missing={incompatible.missing_keys}, unexpected={unexpected}")
    model.eval()
    for parameter in model.parameters():
        parameter.requires_grad_(False)
    return model


def read_sequences(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["sequences"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]
    records: list[dict[str, str]] = []
    for row in rows:
        item = dict(zip(header, row))
        sequence = str(item.get("sequence") or "").strip().upper()
        if sequence:
            records.append(
                {
                    "sample_id": str(item["sequence_id"]),
                    "sequence": sequence,
                    "length": len(sequence),
                    "fold": str(item.get("fold") or ""),
                }
            )
    workbook.close()
    require(records, f"No sequences in {path}")
    return records


def cache_name(sample_id: str) -> str:
    return hashlib.sha256(sample_id.encode("utf-8")).hexdigest() + ".npz"


def main() -> None:
    args = parse_args()
    require(args.stride > 0, "stride must be positive")
    require(args.batch_size > 0, "batch-size must be positive")
    lengths = tuple(sorted({int(x) for x in args.window_lengths.split(",") if x.strip()}))
    require(lengths and all(x > 0 for x in lengths), "invalid window lengths")
    require(args.stride <= lengths[0], "stride cannot exceed the minimum window length")
    records = read_sequences(args.apr_workbook)
    args.output_dir.mkdir(parents=True, exist_ok=False)
    sample_dir = args.output_dir / "samples"
    sample_dir.mkdir()

    device = torch.device(args.device)
    model = load_peptide_teacher(args.backbone, args.peptide_checkpoint, device)

    manifest: list[dict[str, object]] = []
    for record_index, record in enumerate(records, start=1):
        sequence = record["sequence"]
        windows = enumerate_windows(sequence, lengths, args.stride)

        hidden_rows: list[np.ndarray] = []
        probability_rows: list[np.ndarray] = []
        local_mean_rows: list[np.ndarray] = []
        with torch.inference_mode():
            for offset in range(0, len(windows), args.batch_size):
                batch = windows[offset : offset + args.batch_size]
                output = model.forward_sequence(
                    [item[2] for item in batch], isolate_peptide_rng=True
                )
                hidden = output["peptide_hidden"].detach().float().cpu().numpy()
                probabilities = torch.sigmoid(output["peptide_logits"]).detach().float().cpu().numpy()
                local_logits = output.get("aggregation_propensity_logits")
                if local_logits is None:
                    local_mean = probabilities
                else:
                    mask = output["residue_mask"].bool()
                    require(mask.shape == local_logits.shape, "local-logit/mask shape mismatch")
                    require(torch.all(mask.sum(dim=1) > 0).item(), "Empty teacher window")
                    local_mean = ((torch.sigmoid(local_logits).float() * mask).sum(dim=1)
                                  / mask.sum(dim=1)).detach().cpu().numpy()
                for row, (_, window_length, _, _) in enumerate(batch):
                    # A single fixed-width token represents the whole window.
                    hidden_rows.append(hidden[row, :window_length].mean(axis=0).astype(np.float32))
                    probability_rows.append(np.asarray(probabilities[row], dtype=np.float32))
                    local_mean_rows.append(np.asarray(local_mean[row], dtype=np.float32))

        output_path = sample_dir / cache_name(record["sample_id"])
        starts = np.asarray([item[0] for item in windows], dtype=np.int32)
        window_lengths = np.asarray([item[1] for item in windows], dtype=np.int16)
        scale_ids = np.asarray([item[3] for item in windows], dtype=np.int8)
        window_embedding = np.stack(hidden_rows).astype(np.float16)
        window_probability = np.asarray(probability_rows, dtype=np.float32)
        window_local_probability = np.asarray(local_mean_rows, dtype=np.float32)
        covered = np.zeros(len(sequence), dtype=np.int32)
        for start, length in zip(starts.tolist(), window_lengths.tolist()):
            covered[start : start + length] += 1
        require(np.all(covered > 0), f"Window tokens do not cover all residues: {record['sample_id']}")
        np.savez_compressed(
            output_path,
            sample_id=np.asarray(record["sample_id"]),
            sequence=np.asarray(sequence),
            window_start=starts,
            window_length=window_lengths,
            window_scale_id=scale_ids,
            window_embedding=window_embedding,
            window_probability=window_probability,
            window_local_probability=window_local_probability,
        )
        manifest.append({**record, "cache_file": output_path.name, "window_lengths": list(lengths), "stride": args.stride, "window_count": len(windows)})
        print(json.dumps({"status": "window_cache", "completed": record_index, "total": len(records), "sample_id": record["sample_id"], "window_count": len(windows)}), flush=True)

    metadata = {
        "definition": "Independent Aggregation-aware Peptide tokens for APR transfer",
        "peptide_checkpoint": str(args.peptide_checkpoint),
        "peptide_checkpoint_sha256": hashlib.sha256(args.peptide_checkpoint.read_bytes()).hexdigest(),
        "window_lengths": list(lengths),
        "stride": args.stride,
        "batch_size": args.batch_size,
        "local_probability_policy": MASKED_LOCAL_MEAN,
        "short_sequence_policy": "full sequence token in smallest scale slot with true length",
        "tail_window_policy": "right-aligned final window added when regular stride misses the tail",
        "token_width": 320,
        "token_definition": "mean peptide_hidden over each individual sliding window",
        "sample_count": len(records),
        "records": manifest,
    }
    (args.output_dir / "manifest.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "complete", "output_dir": str(args.output_dir)}), flush=True)


if __name__ == "__main__":
    main()
