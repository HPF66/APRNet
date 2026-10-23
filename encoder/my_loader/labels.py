from __future__ import annotations
import argparse
import csv
import json
import random
import sys
from dataclasses import dataclass
from pathlib import Path
import numpy as np
import torch
from openpyxl import load_workbook
from sklearn.metrics import accuracy_score, average_precision_score, balanced_accuracy_score, confusion_matrix, f1_score, matthews_corrcoef, roc_auc_score
from torch import Tensor, nn
from torch.nn import functional as F
from torch.utils.data import DataLoader, Dataset
import encoder.my_loader.window_data as trainer
HIDDEN_SIZE = 192
EVIDENCE_FEATURES = 24
WINDOW_SCALES = (6, 8, 10, 12, 16, 20)

def require(ok: bool, message: str) -> None:
    if not ok:
        raise ValueError(message)

def read_fold_map(path: Path) -> dict[str, int]:
    with path.open(newline='', encoding='utf-8') as handle:
        rows = list(csv.DictReader(handle))
    result = {row['item_id']: int(row['fold']) for row in rows if row.get('dataset') == 'apr'}
    require(result, f'No APR folds in {path}')
    return result

def read_proteins(path: Path, fold_csv: Path) -> list[tuple[str, str, np.ndarray, int]]:
    fold_map = read_fold_map(fold_csv)
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook['sequences'].iter_rows(values_only=True)
    header = [str(v) for v in next(rows)]
    result: list[tuple[str, str, np.ndarray, int]] = []
    for values in rows:
        item = dict(zip(header, values))
        sample_id = str(item['sequence_id'])
        sequence = str(item.get('sequence') or '').strip().upper()
        if not sequence:
            continue
        require(sample_id in fold_map, f'Missing fold for {sample_id}')
        mask = str(item.get('apr_positive_mask') or '').strip()
        require(len(mask) == len(sequence), f'APR mask mismatch for {sample_id}')
        target = np.fromiter((1 if char == '1' else 0 for char in mask), dtype=np.float32, count=len(sequence))
        result.append((sample_id, sequence, target, fold_map[sample_id]))
    workbook.close()
    require(result, 'No APR proteins')
    return result
