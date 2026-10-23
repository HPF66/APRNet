from __future__ import annotations

import json
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Sequence

import torch
from torch.utils.data import Dataset, Sampler


AMINO_ACIDS = "ACDEFGHIKLMNPQRSTVWY"
AMINO_TO_INDEX = {aa: i for i, aa in enumerate(AMINO_ACIDS)}
UNKNOWN_INDEX = len(AMINO_ACIDS)
STRUCTURE_GEOMETRIC_FEATURE_DIM = 9


def _read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    """Read one cleaned workbook sheet without loading the whole workbook into pandas."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{path} does not contain sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    records = [dict(zip(headers, row)) for row in rows]
    workbook.close()
    return records


def _as_sequence(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).upper()


def parse_target_chain_ids(value: Any) -> list[str]:
    if value is None:
        return []
    return [item for item in re.split(r"[;,|\s]+", str(value).strip()) if item]


def parse_coordinate_residue_keys(value: Any) -> tuple[tuple[int, str], ...]:
    """Parse cleaned PDB residue keys such as ``42`` or ``42A``."""
    if value is None:
        return ()
    keys: list[tuple[int, str]] = []
    for token in str(value).split(","):
        match = re.fullmatch(r"\s*(-?\d+)([A-Za-z]?)\s*", token)
        if match is None:
            raise ValueError(f"Invalid coordinate residue key: {token!r}")
        keys.append((int(match.group(1)), match.group(2).upper()))
    return tuple(keys)


def encode_apr_mask(mask: Any, length: int) -> torch.Tensor:
    """Encode an APR mask as 1, 0, or -100 (ignored)."""
    text = "" if mask is None else str(mask).strip()
    target = torch.full((length,), -100.0, dtype=torch.float32)
    for index, char in enumerate(text[:length]):
        if char == "1":
            target[index] = 1.0
        elif char == "0":
            target[index] = 0.0
    return target


def _negative_intervals_from_aggregating(path: Path) -> dict[str, list[tuple[int, int]]]:
    """Return strict non-amyloid intervals keyed by their exact protein sequence."""
    protein_rows = _read_sheet(path, "protein_sequences")
    sequence_by_id = {
        str(row["sequence_id"]): _as_sequence(row["sequence"])
        for row in protein_rows
        if row.get("sequence_id") and row.get("sequence")
    }
    intervals: dict[str, list[tuple[int, int]]] = {}
    for row in _read_sheet(path, "mapped_regions"):
        if int(row.get("label", -1)) != 0:
            continue
        confidence = str(row.get("mapping_confidence") or "").casefold()
        overlap = row.get("opposite_label_overlap")
        if confidence not in {"exact", "constructed_variant", "mutant_constructed"}:
            continue
        if overlap not in (None, 0, "0"):
            continue
        sequence = sequence_by_id.get(str(row.get("sequence_id")))
        if not sequence:
            continue
        start = int(row["start"])
        end = int(row["end"])
        if 1 <= start <= end <= len(sequence):
            intervals.setdefault(sequence, []).append((start, end))
    return intervals


@dataclass(frozen=True)
class PeptideRecord:
    peptide_id: str
    sequence: str
    label: int
    protein_id: str | None
    split_group_id: str | None = None


@dataclass(frozen=True)
class APRRecord:
    sequence_id: str
    sequence: str
    target: torch.Tensor
    label_weight: torch.Tensor
    training_only: bool = False
    source_mixed: bool = False


@dataclass(frozen=True)
class StructureRecord:
    sample_id: str
    pdb_id: str
    label: int
    structure_file: str
    target_chain_ids: tuple[str, ...]
    chain_sequences: tuple[str, ...]
    chain_residue_keys: tuple[tuple[tuple[int, str], ...], ...]
    split_group_id: str | None


class AggregatingPeptideDataset(Dataset[PeptideRecord]):
    def __init__(self, workbook: str | Path):
        self.records = [
            PeptideRecord(
                peptide_id=str(row["peptide_id"]),
                sequence=_as_sequence(row["peptide_sequence"]),
                label=int(row["label"]),
                protein_id=(str(row["protein_id"]) if row.get("protein_id") else None),
                split_group_id=(
                    str(row["protein_split_group_id"])
                    if row.get("protein_split_group_id")
                    else None
                ),
            )
            for row in _read_sheet(Path(workbook), "peptides")
            if row.get("peptide_sequence")
        ]

    def __len__(self) -> int:
        return len(self.records)

    def __getitem__(self, index: int) -> PeptideRecord:
        return self.records[index]


class APRDataset(Dataset[APRRecord]):
    def __init__(
        self,
        workbook: str | Path,
        aggregating_workbook: str | Path | None = None,
        use_aggregating_negatives: bool = True,
        weak_negative_workbook: str | Path | None = None,
        weak_negative_weight: float = 0.25,
        mapped_negative_weight: float = 1.0,
    ):
        if not 0 < weak_negative_weight <= 1:
            raise ValueError("weak_negative_weight must be in (0, 1]")
        if mapped_negative_weight <= 0:
            raise ValueError("mapped_negative_weight must be positive")
        apr_rows = _read_sheet(Path(workbook), "sequences")
        materialized_negatives = any(
            "0" in str(row.get("apr_positive_mask") or "") for row in apr_rows
        )
        negative_intervals: dict[str, list[tuple[int, int]]] = {}
        if use_aggregating_negatives and aggregating_workbook and not materialized_negatives:
            negative_intervals = _negative_intervals_from_aggregating(
                Path(aggregating_workbook)
            )

        records: list[APRRecord] = []
        for row in apr_rows:
            sequence = _as_sequence(row["sequence"])
            if not sequence:
                continue
            target = encode_apr_mask(row.get("apr_positive_mask"), len(sequence))
            # Only strict, non-overlapping experimental non-amyloid intervals
            # are converted to zero labels. Everything else stays ignored.
            for start, end in negative_intervals.get(sequence, []):
                interval = slice(start - 1, end)
                existing = target[interval]
                conflict = existing == 1
                target[interval] = torch.where(
                    conflict,
                    torch.full_like(existing, -100.0),
                    torch.zeros_like(existing),
                )
            label_weight = (target != -100).to(dtype=torch.float32)
            label_sources = {
                item.strip()
                for item in str(row.get("label_source") or "").split(";")
                if item.strip()
            }
            if "aggregating_experimental_negative" in label_sources:
                label_weight[target == 0] *= float(mapped_negative_weight)
            records.append(
                APRRecord(
                    sequence_id=str(row["sequence_id"]),
                    sequence=sequence,
                    target=target,
                    label_weight=label_weight,
                    source_mixed=bool(
                        torch.any(target == 1).item()
                        and torch.any(target == 0).item()
                    ),
                )
            )

        if weak_negative_workbook is not None:
            weak_path = Path(weak_negative_workbook)
            for row in _read_sheet(weak_path, "weak_negative_windows"):
                sequence = _as_sequence(row.get("context_sequence"))
                if not sequence:
                    raise ValueError(
                        f"Empty weak-negative context for {row.get('weak_sequence_id')}"
                    )
                target = encode_apr_mask(row.get("weak_label_mask"), len(sequence))
                if torch.any(target == 1) or not torch.any(target == 0):
                    raise ValueError(
                        f"Invalid weak-negative mask for {row.get('weak_sequence_id')}"
                    )
                label_weight = torch.zeros(len(sequence), dtype=torch.float32)
                label_weight[target == 0] = float(weak_negative_weight)
                records.append(
                    APRRecord(
                        sequence_id=str(row["weak_sequence_id"]),
                        sequence=sequence,
                        target=target,
                        label_weight=label_weight,
                        training_only=True,
                    )
                )
        identifiers = [record.sequence_id for record in records]
        if len(identifiers) != len(set(identifiers)):
            raise ValueError("APR sequence IDs must be unique")
        self.records = records

    def __len__(self) -> int:
        return len(self.records)

    def __getitem__(self, index: int) -> APRRecord:
        return self.records[index]


class PDBGraphBuilder:
    """Build a residue graph from the selected PDB chains."""

    STANDARD_THREE_TO_ONE = {
        "ALA": "A", "ARG": "R", "ASN": "N", "ASP": "D", "CYS": "C",
        "GLN": "Q", "GLU": "E", "GLY": "G", "HIS": "H", "ILE": "I",
        "LEU": "L", "LYS": "K", "MET": "M", "PHE": "F", "PRO": "P",
        "SER": "S", "THR": "T", "TRP": "W", "TYR": "Y", "VAL": "V",
    }

    def __init__(
        self,
        residue_code_map: dict[str, str] | None = None,
        cutoff: float = 8.0,
        edge_weighting: str = "binary",
        distance_scale: float = 4.0,
        normalize_position: bool = False,
        geometric_features: bool = False,
    ):
        try:
            from Bio.PDB import PDBParser
        except ImportError as exc:  # pragma: no cover - exercised on server
            raise RuntimeError("Biopython is required for PDB feature extraction") from exc
        self.parser = PDBParser(QUIET=True)
        self.residue_code_map = {
            **self.STANDARD_THREE_TO_ONE,
            **(residue_code_map or {}),
        }
        if cutoff <= 0:
            raise ValueError("PDB graph cutoff must be positive")
        if edge_weighting not in {"binary", "gaussian", "inverse"}:
            raise ValueError(f"Unknown edge weighting: {edge_weighting}")
        if distance_scale <= 0:
            raise ValueError("PDB graph distance scale must be positive")
        self.cutoff = cutoff
        self.edge_weighting = edge_weighting
        self.distance_scale = distance_scale
        self.normalize_position = normalize_position
        self.geometric_features = geometric_features

    @staticmethod
    def _geometric_node_features(
        coordinates: torch.Tensor,
        distance: torch.Tensor,
        chain_indices: list[int],
        sequence_positions: list[int],
    ) -> torch.Tensor:
        count = coordinates.shape[0]
        device = coordinates.device
        chain = torch.tensor(chain_indices, dtype=torch.long, device=device)
        position = torch.tensor(sequence_positions, dtype=torch.long, device=device)
        same_chain = chain.unsqueeze(0) == chain.unsqueeze(1)
        sequence_separation = (position.unsqueeze(0) - position.unsqueeze(1)).abs()
        non_self = ~torch.eye(count, dtype=torch.bool, device=device)
        normalization = torch.log1p(
            torch.tensor(float(max(count, 2)), dtype=coordinates.dtype, device=device)
        )

        features = torch.zeros(
            count,
            STRUCTURE_GEOMETRIC_FEATURE_DIM,
            dtype=coordinates.dtype,
            device=device,
        )
        for column, cutoff in enumerate((6.0, 8.0, 10.0)):
            neighbors = ((distance <= cutoff) & non_self).sum(dim=1).to(coordinates.dtype)
            features[:, column] = torch.log1p(neighbors) / normalization

        long_range = (
            (distance <= 8.0)
            & same_chain
            & (sequence_separation >= 8)
            & non_self
        )
        inter_chain = (distance <= 8.0) & ~same_chain
        features[:, 3] = torch.log1p(long_range.sum(dim=1).to(coordinates.dtype)) / normalization
        features[:, 4] = torch.log1p(inter_chain.sum(dim=1).to(coordinates.dtype)) / normalization

        center = coordinates.mean(dim=0, keepdim=True)
        radial_distance = torch.linalg.vector_norm(coordinates - center, dim=1)
        radius_of_gyration = torch.sqrt(radial_distance.pow(2).mean()).clamp_min(1e-6)
        features[:, 5] = (radial_distance / radius_of_gyration).clamp_max(4.0) / 4.0

        for index in range(1, count - 1):
            if chain_indices[index - 1] != chain_indices[index] or chain_indices[index + 1] != chain_indices[index]:
                continue
            left = coordinates[index - 1] - coordinates[index]
            right = coordinates[index + 1] - coordinates[index]
            denominator = (
                torch.linalg.vector_norm(left) * torch.linalg.vector_norm(right)
            ).clamp_min(1e-6)
            features[index, 6] = torch.dot(left, right) / denominator

        for index in range(count - 3):
            if len(set(chain_indices[index : index + 4])) != 1:
                continue
            p0, p1, p2, p3 = coordinates[index : index + 4]
            b0 = p0 - p1
            b1 = p2 - p1
            b2 = p3 - p2
            b1_unit = b1 / torch.linalg.vector_norm(b1).clamp_min(1e-6)
            v = b0 - torch.dot(b0, b1_unit) * b1_unit
            w = b2 - torch.dot(b2, b1_unit) * b1_unit
            x = torch.dot(v, w)
            y = torch.dot(torch.cross(b1_unit, v, dim=0), w)
            angle = torch.atan2(y, x)
            features[index + 1 : index + 3, 7] = torch.sin(angle)
            features[index + 1 : index + 3, 8] = torch.cos(angle)
        return features

    def _residue_letter(self, residue: Any) -> str | None:
        code = str(residue.resname).strip().upper()
        letter = self.residue_code_map.get(code)
        if letter in AMINO_TO_INDEX or letter == "X":
            return letter
        return None

    @staticmethod
    def _residue_coordinate(residue: Any) -> tuple[Any, bool]:
        """Use CA when present; modified residues fall back to their heavy-atom centroid."""
        if "CA" in residue:
            return residue["CA"].coord.astype("float32"), True
        heavy_atoms = [
            atom.coord.astype("float32")
            for atom in residue.get_atoms()
            if str(getattr(atom, "element", "")).strip().upper() != "H"
        ]
        if not heavy_atoms:
            raise ValueError(
                f"No usable atomic coordinates for {residue.resname} residue {residue.id}"
            )
        return sum(heavy_atoms) / len(heavy_atoms), False

    def build(
        self,
        pdb_path: Path,
        chain_ids: Sequence[str],
        chain_residue_keys: Sequence[Sequence[tuple[int, str]]] | None = None,
    ) -> tuple[torch.Tensor, torch.Tensor, torch.Tensor]:
        structure = self.parser.get_structure(pdb_path.stem, str(pdb_path))
        model = next(structure.get_models())
        selected = list(chain_ids) if chain_ids else [chain.id for chain in model]
        if chain_residue_keys is not None and len(chain_residue_keys) != len(selected):
            raise ValueError(
                f"Residue-key chain count mismatch for {pdb_path}: "
                f"keys={len(chain_residue_keys)}, chains={len(selected)}"
            )
        nodes: list[list[float]] = []
        coordinates: list[list[float]] = []
        chain_indices: list[int] = []
        sequence_positions: list[int] = []
        previous_by_chain: dict[str, int] = {}

        for chain_index, chain_id in enumerate(selected):
            if chain_id not in model:
                raise ValueError(f"Selected chain {chain_id!r} is absent from {pdb_path}")
            chain = model[chain_id]
            requested_keys = (
                chain_residue_keys[chain_index] if chain_residue_keys is not None else None
            )
            if requested_keys is None:
                residues = list(chain)
            else:
                by_key = {
                    (int(residue.id[1]), str(residue.id[2]).strip().upper()): residue
                    for residue in chain
                }
                residues = []
                for key in requested_keys:
                    residue = by_key.get(key)
                    if residue is None:
                        raise ValueError(
                            f"Coordinate residue {chain_id}:{key[0]}{key[1]} is absent from {pdb_path}"
                        )
                    residues.append(residue)
            position = 0
            for residue in residues:
                letter = self._residue_letter(residue)
                if letter is None:
                    raise ValueError(
                        f"Unsupported coordinate residue {residue.resname} {residue.id} in {pdb_path}"
                    )
                coord, has_ca = self._residue_coordinate(residue)
                aa_index = AMINO_TO_INDEX.get(letter, UNKNOWN_INDEX)
                feature = [0.0] * (len(AMINO_ACIDS) + 2)
                feature[aa_index] = 1.0
                feature[-2] = float(has_ca)
                feature[-1] = (
                    float(position) / max(len(residues) - 1, 1)
                    if self.normalize_position
                    else float(position) / max(position + 1, 1)
                )
                nodes.append(feature)
                coordinates.append(coord.tolist())
                chain_indices.append(chain_index)
                sequence_positions.append(position)
                previous_by_chain[chain_id] = len(nodes) - 1
                position += 1

        if not nodes:
            raise ValueError(f"No supported CA residues found in {pdb_path}")

        node_features = torch.tensor(nodes, dtype=torch.float32)
        coords = torch.tensor(coordinates, dtype=torch.float32)
        distance = torch.cdist(coords, coords)
        if self.geometric_features:
            node_features = torch.cat(
                [
                    node_features,
                    self._geometric_node_features(
                        coords,
                        distance,
                        chain_indices,
                        sequence_positions,
                    ),
                ],
                dim=1,
            )
        contact_mask = distance <= self.cutoff
        if self.edge_weighting == "binary":
            adjacency = contact_mask.to(torch.float32)
        elif self.edge_weighting == "gaussian":
            adjacency = torch.exp(
                -0.5 * (distance / self.distance_scale).pow(2)
            ) * contact_mask.to(torch.float32)
        else:
            adjacency = contact_mask.to(torch.float32) / (1.0 + distance)
        for left, right in zip(range(len(nodes) - 1), range(1, len(nodes))):
            if chain_indices[left] == chain_indices[right]:
                adjacency[left, right] = 1.0
                adjacency[right, left] = 1.0
        adjacency.fill_diagonal_(0.0)
        node_mask = torch.ones(len(nodes), dtype=torch.bool)
        return node_features, adjacency, node_mask


class AmyloidStructureDataset(Dataset[StructureRecord]):
    def __init__(
        self,
        workbook: str | Path,
        pdb_dir: str | Path,
        graph_cutoff: float = 8.0,
        graph_edge_weighting: str = "binary",
        graph_distance_scale: float = 4.0,
        normalize_position: bool = False,
        geometric_features: bool = False,
        peptide_evidence_file: str | Path | None = None,
    ):
        workbook = Path(workbook)
        self.pdb_dir = Path(pdb_dir)
        residue_code_map = {
            str(row["residue_code"]).strip().upper(): str(row["sequence_token"]).strip().upper()
            for row in _read_sheet(workbook, "residue_codes")
            if row.get("residue_code") and row.get("sequence_token")
        }
        chains_by_sample: dict[str, list[tuple[int, str, str, tuple[tuple[int, str], ...]]]] = {}
        for row in _read_sheet(workbook, "chains"):
            sample_id = str(row["sample_id"])
            chain_id = str(row["chain_id"])
            sequence = _as_sequence(row.get("coordinate_sequence"))
            if not sequence:
                raise ValueError(f"Empty coordinate sequence for {sample_id} chain {chain_id}")
            expected_length = int(row.get("coordinate_length") or len(sequence))
            if len(sequence) != expected_length:
                raise ValueError(
                    f"Coordinate sequence length mismatch for {sample_id} chain {chain_id}"
                )
            residue_keys = parse_coordinate_residue_keys(row.get("coordinate_residue_keys"))
            if len(residue_keys) != len(sequence):
                raise ValueError(
                    f"Coordinate residue-key length mismatch for {sample_id} chain {chain_id}"
                )
            chains_by_sample.setdefault(sample_id, []).append(
                (int(row["chain_index"]), chain_id, sequence, residue_keys)
            )

        records: list[StructureRecord] = []
        for row in _read_sheet(workbook, "samples"):
            sample_id = str(row["sample_id"])
            target_chain_ids = tuple(parse_target_chain_ids(row.get("target_chain_ids")))
            chain_rows = sorted(chains_by_sample.get(sample_id, []))
            chain_ids = tuple(item[1] for item in chain_rows)
            if not chain_rows:
                raise ValueError(f"No coordinate chains found for structure sample {sample_id}")
            if target_chain_ids and chain_ids != target_chain_ids:
                raise ValueError(
                    f"Chain order mismatch for {sample_id}: samples={target_chain_ids}, "
                    f"chains={chain_ids}"
                )
            records.append(
                StructureRecord(
                    sample_id=sample_id,
                    pdb_id=str(row["pdb_id"]),
                    label=int(row["label"]),
                    structure_file=str(row["structure_file"]),
                    target_chain_ids=target_chain_ids,
                    chain_sequences=tuple(item[2] for item in chain_rows),
                    chain_residue_keys=tuple(item[3] for item in chain_rows),
                    split_group_id=(
                        str(row["exact_sequence_split_group_id"])
                        if row.get("exact_sequence_split_group_id")
                        else None
                    ),
                )
            )
        self.records = records
        self.peptide_evidence: dict[str, torch.Tensor] = {}
        if peptide_evidence_file is not None:
            evidence_path = Path(peptide_evidence_file)
            payload = json.loads(evidence_path.read_text(encoding="utf-8"))
            samples = payload.get("samples")
            if not isinstance(samples, dict):
                raise ValueError(
                    f"Peptide evidence file has no sample mapping: {evidence_path}"
                )
            for record in records:
                sample = samples.get(record.sample_id)
                if not isinstance(sample, dict):
                    raise ValueError(
                        f"Missing Peptide evidence for structure {record.sample_id}"
                    )
                expected_chain_lengths = [
                    len(sequence) for sequence in record.chain_sequences
                ]
                if sample.get("chain_lengths") != expected_chain_lengths:
                    raise ValueError(
                        f"Peptide evidence chain mismatch for {record.sample_id}"
                    )
                evidence = torch.tensor(
                    sample.get("residue_evidence", []), dtype=torch.float32
                )
                expected_length = sum(expected_chain_lengths)
                if evidence.numel() != expected_length:
                    raise ValueError(
                        f"Peptide evidence length mismatch for {record.sample_id}: "
                        f"evidence={evidence.numel()}, sequence={expected_length}"
                    )
                if not torch.all(torch.isfinite(evidence)) or not torch.all(
                    (evidence >= 0) & (evidence <= 1)
                ):
                    raise ValueError(
                        f"Invalid Peptide evidence probabilities for {record.sample_id}"
                    )
                self.peptide_evidence[record.sample_id] = evidence
        self._builder = PDBGraphBuilder(
            residue_code_map=residue_code_map,
            cutoff=graph_cutoff,
            edge_weighting=graph_edge_weighting,
            distance_scale=graph_distance_scale,
            normalize_position=normalize_position,
            geometric_features=geometric_features,
        )

    def __len__(self) -> int:
        return len(self.records)

    # Cover the full 561-sample dataset so cross-validation epochs do not rebuild PDB graphs.
    @lru_cache(maxsize=1024)
    def _graph(
        self,
        structure_file: str,
        chain_ids: tuple[str, ...],
        chain_residue_keys: tuple[tuple[tuple[int, str], ...], ...],
    ):
        return self._builder.build(
            self.pdb_dir / structure_file,
            chain_ids,
            chain_residue_keys,
        )

    def __getitem__(self, index: int) -> dict[str, Any]:
        record = self.records[index]
        node_features, adjacency, node_mask = self._graph(
            record.structure_file,
            record.target_chain_ids,
            record.chain_residue_keys,
        )
        sequence_length = sum(len(sequence) for sequence in record.chain_sequences)
        if sequence_length != node_features.shape[0]:
            raise ValueError(
                f"PDB/coordinate-sequence residue mismatch for {record.sample_id}: "
                f"graph={node_features.shape[0]}, sequences={sequence_length}"
            )
        return {
            "sample_id": record.sample_id,
            "label": torch.tensor(float(record.label), dtype=torch.float32),
            "group_id": record.split_group_id,
            "chain_sequences": record.chain_sequences,
            "node_features": node_features,
            "adjacency": adjacency,
            "node_mask": node_mask,
            "peptide_evidence": self.peptide_evidence.get(record.sample_id),
        }


class BalancedBinaryBatchSampler(Sampler[list[int]]):
    """Yield class-balanced batches while keeping epoch length unchanged."""

    def __init__(self, labels: Sequence[int], batch_size: int):
        if batch_size < 2 or batch_size % 2:
            raise ValueError("Balanced binary batches require an even batch size")
        self.positive = [index for index, label in enumerate(labels) if int(label) == 1]
        self.negative = [index for index, label in enumerate(labels) if int(label) == 0]
        if not self.positive or not self.negative:
            raise ValueError("Balanced binary batches require both classes")
        self.batch_size = batch_size
        self.half_batch = batch_size // 2
        self.num_batches = (len(labels) + batch_size - 1) // batch_size

    @staticmethod
    def _sample(pool: list[int], count: int) -> list[int]:
        output: list[int] = []
        while len(output) < count:
            order = torch.randperm(len(pool)).tolist()
            output.extend(pool[index] for index in order)
        return output[:count]

    def __iter__(self):
        count_per_class = self.num_batches * self.half_batch
        positive = self._sample(self.positive, count_per_class)
        negative = self._sample(self.negative, count_per_class)
        for batch_index in range(self.num_batches):
            start = batch_index * self.half_batch
            end = start + self.half_batch
            batch = positive[start:end] + negative[start:end]
            order = torch.randperm(len(batch)).tolist()
            yield [batch[index] for index in order]

    def __len__(self) -> int:
        return self.num_batches


def collate_peptides(records: Iterable[PeptideRecord]) -> dict[str, Any]:
    records = list(records)
    return {
        "ids": [record.peptide_id for record in records],
        "sequences": [record.sequence for record in records],
        "labels": torch.tensor([record.label for record in records], dtype=torch.float32),
    }


def apr_training_windows(
    records: Iterable[APRRecord], flank: int, label_weight_scale: float = 1.0
) -> list[APRRecord]:
    if flank < 0:
        raise ValueError("APR training-window flank must be non-negative")
    if label_weight_scale < 0:
        raise ValueError("APR training-window label-weight scale must be non-negative")
    windows: list[APRRecord] = []
    for record in records:
        valid_indices = torch.nonzero(record.target != -100, as_tuple=False).flatten()
        if valid_indices.numel() == 0:
            continue
        run_starts = [int(valid_indices[0])]
        run_ends: list[int] = []
        for previous, current in zip(valid_indices[:-1], valid_indices[1:]):
            if int(current) != int(previous) + 1:
                run_ends.append(int(previous) + 1)
                run_starts.append(int(current))
        run_ends.append(int(valid_indices[-1]) + 1)

        for run_index, (run_start, run_end) in enumerate(
            zip(run_starts, run_ends), start=1
        ):
            window_start = max(0, run_start - flank)
            window_end = min(len(record.sequence), run_end + flank)
            target = torch.full(
                (window_end - window_start,), -100.0, dtype=record.target.dtype
            )
            label_weight = torch.zeros(
                window_end - window_start, dtype=record.label_weight.dtype
            )
            local_start = run_start - window_start
            local_end = run_end - window_start
            target[local_start:local_end] = record.target[run_start:run_end]
            label_weight[local_start:local_end] = (
                label_weight_scale * record.label_weight[run_start:run_end]
            )
            windows.append(
                APRRecord(
                    sequence_id=f"{record.sequence_id}:window{run_index}",
                    sequence=record.sequence[window_start:window_end],
                    target=target,
                    label_weight=label_weight,
                    training_only=record.training_only,
                    source_mixed=record.source_mixed,
                )
            )
    return windows


def collate_apr(
    records: Iterable[APRRecord],
    training_window_flank: int | None = None,
    training_window_full_mix: float = 0.0,
) -> dict[str, Any]:
    records = list(records)
    if training_window_flank is not None:
        if not 0.0 <= training_window_full_mix <= 1.0:
            raise ValueError("APR training-window full mix must be in [0, 1]")
        windows = apr_training_windows(
            records,
            training_window_flank,
            label_weight_scale=1.0 - training_window_full_mix,
        )
        if training_window_full_mix > 0:
            full_records = [
                APRRecord(
                    sequence_id=f"{record.sequence_id}:full",
                    sequence=record.sequence,
                    target=record.target,
                    label_weight=training_window_full_mix * record.label_weight,
                    training_only=record.training_only,
                    source_mixed=record.source_mixed,
                )
                for record in records
            ]
            records = full_records + windows
        else:
            records = windows
    if not records:
        raise ValueError("APR batch contains no labeled records")
    max_length = max(len(record.sequence) for record in records)
    targets = torch.full((len(records), max_length), -100.0, dtype=torch.float32)
    label_weights = torch.zeros((len(records), max_length), dtype=torch.float32)
    for row, record in enumerate(records):
        targets[row, : len(record.target)] = record.target
        label_weights[row, : len(record.label_weight)] = record.label_weight
    return {
        "ids": [record.sequence_id for record in records],
        "sequences": [record.sequence for record in records],
        "targets": targets,
        "label_weights": label_weights,
        "training_only": torch.tensor(
            [record.training_only for record in records], dtype=torch.bool
        ),
        "source_mixed": torch.tensor(
            [record.source_mixed for record in records], dtype=torch.bool
        ),
    }


def collate_structures(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    records = list(records)
    max_nodes = max(item["node_features"].shape[0] for item in records)
    feature_dim = records[0]["node_features"].shape[1]
    node_features = torch.zeros(len(records), max_nodes, feature_dim)
    adjacency = torch.zeros(len(records), max_nodes, max_nodes)
    node_mask = torch.zeros(len(records), max_nodes, dtype=torch.bool)
    evidence_present = [item["peptide_evidence"] is not None for item in records]
    if any(evidence_present) and not all(evidence_present):
        raise ValueError("Structure batch mixes samples with and without Peptide evidence")
    peptide_evidence = (
        torch.zeros(len(records), max_nodes) if all(evidence_present) else None
    )
    for index, item in enumerate(records):
        count = item["node_features"].shape[0]
        node_features[index, :count] = item["node_features"]
        adjacency[index, :count, :count] = item["adjacency"]
        node_mask[index, :count] = item["node_mask"]
        if peptide_evidence is not None:
            evidence = item["peptide_evidence"]
            if evidence.numel() != count:
                raise ValueError(
                    f"Structure/Peptide evidence mismatch for {item['sample_id']}: "
                    f"graph={count}, evidence={evidence.numel()}"
                )
            peptide_evidence[index, :count] = evidence
    return {
        "ids": [item["sample_id"] for item in records],
        "labels": torch.stack([item["label"] for item in records]),
        "chain_sequences": [item["chain_sequences"] for item in records],
        "node_features": node_features,
        "adjacency": adjacency,
        "node_mask": node_mask,
        "peptide_evidence": peptide_evidence,
    }
