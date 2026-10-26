#!/usr/bin/env python3
"""Export frozen Boundary-Aware Context V3 embeddings for unique sequences.

This operates on the unique sequence records in the supplied workbook.  PDB
sample-level chain replication/concatenation is a separate downstream step.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--workbook", type=Path, required=True)
    p.add_argument("--window-cache", type=Path, required=True)
    p.add_argument("--esm-cache", type=Path, required=True)
    p.add_argument("--checkpoint", type=Path, default=Path(__file__).resolve().parents[2] / "pretrained/encoder_epoch20.pt")
    p.add_argument("--output-dir", type=Path, required=True)
    p.add_argument("--device", default="cuda")
    p.add_argument("--batch-size", type=int, default=4)
    p.add_argument("--max-length", type=int, default=512)
    p.add_argument("--overlap", type=int, default=64)
    p.add_argument("--seed", type=int, default=2026)
    return p.parse_args()


def read_sequences(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["sequences"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(v) for v in next(rows)]
    records = []
    for row in rows:
        item = dict(zip(header, row))
        sid = str(item.get("sequence_id") or "").strip()
        sequence = str(item.get("sequence") or "").strip().upper()
        if sid and sequence:
            records.append((sid, sequence))
    workbook.close()
    require(records, f"No sequences found in {path}")
    require(len({sid for sid, _ in records}) == len(records), "Duplicate sequence_id")
    return records


def main() -> None:
    args = parse_args()
    require(args.batch_size > 0, "batch-size must be positive")
    require(args.max_length > 0 and 0 <= args.overlap < args.max_length, "invalid chunk geometry")
    import encoder.my_loader.window_data as trainer
    import encoder.model.residue_encoder as boundary_v3

    records = read_sequences(args.workbook)
    proteins = [
        trainer.Protein(sid, sequence, np.zeros(len(sequence), dtype=np.int8), 0)
        for sid, sequence in records
    ]
    protein_by_id = {p.sample_id: p for p in proteins}
    device = torch.device(args.device)
    torch.manual_seed(args.seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(args.seed)

    model, payload = boundary_v3.load_encoder_checkpoint(args.checkpoint, device, frozen=True)
    from encoder.utils.lineage import verify_encoder_inputs
    input_contract = verify_encoder_inputs(args.checkpoint, payload, args.window_cache, args.esm_cache)
    store = trainer.load_store(proteins, args.window_cache, args.esm_cache)
    require(all(np.count_nonzero(item['esm']) == 0 for item in store.values()), "Nonzero base inputs for ESM-free encoder")
    model.eval()
    chunks = trainer.make_chunks(proteins, args.max_length, args.overlap)
    loader = trainer.make_loader(
        chunks, protein_by_id, store, args.batch_size, False, args.seed, 0
    )

    sums = {}
    counts = {}
    with torch.inference_mode():
        for batch_index, raw in enumerate(loader, start=1):
            batch = trainer.move(raw, device)
            out = model(
                batch["base"], batch["tokens"], batch["token_starts"],
                batch["token_lengths"], batch["token_mask"], batch["mask"],
                batch["token_probability"], batch["token_local_probability"],
                batch["token_scale"], batch["residue_evidence"],
                profile_hidden=batch["profile_hidden"],
                profile_local_probability=batch["profile_local_probability"],
                return_encoder=True,
            )
            values = {
                "hapr": out["residue_embedding"].float().cpu().numpy(),
                "pAPR": out["apr_probability"].float().cpu().numpy(),
                "rAPR": out["reliability"].float().cpu().numpy(),
            }
            for row, sid in enumerate(batch["sample_ids"]):
                n = int(raw["mask"][row].sum())
                start = int(batch["starts"][row])
                end = start + n
                length = len(protein_by_id[sid].sequence)
                if sid not in sums:
                    sums[sid] = {
                        "hapr": np.zeros((length, 192), dtype=np.float64),
                        "pAPR": np.zeros(length, dtype=np.float64),
                        "rAPR": np.zeros(length, dtype=np.float64),
                    }
                    counts[sid] = np.zeros(length, dtype=np.int32)
                sums[sid]["hapr"][start:end] += values["hapr"][row, :n]
                sums[sid]["pAPR"][start:end] += values["pAPR"][row, :n]
                sums[sid]["rAPR"][start:end] += values["rAPR"][row, :n]
                counts[sid][start:end] += 1
            if batch_index % 25 == 0 or batch_index == len(loader):
                print(json.dumps({"event": "inference", "batch": batch_index,
                                  "batches": len(loader)}), flush=True)

    require(set(sums) == set(protein_by_id), "Missing unique sequence embedding")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    sample_dir = args.output_dir / "sequences"
    sample_dir.mkdir(exist_ok=True)
    manifest_records = []
    for sid, sequence in records:
        require(np.all(counts[sid] > 0), f"Uncovered residues: {sid}")
        hapr = (sums[sid]["hapr"] / counts[sid][:, None]).astype(np.float32)
        papr = (sums[sid]["pAPR"] / counts[sid]).astype(np.float32)
        rapr = (sums[sid]["rAPR"] / counts[sid]).astype(np.float32)
        require(hapr.shape == (len(sequence), 192), f"HAPR shape mismatch: {sid}")
        require(np.isfinite(hapr).all() and np.isfinite(papr).all() and np.isfinite(rapr).all(),
                f"Nonfinite output: {sid}")
        filename = hashlib.sha256(sid.encode("utf-8")).hexdigest() + ".npz"
        np.savez_compressed(
            sample_dir / filename,
            sequence=np.asarray(sequence),
            sequence_id=np.asarray(sid),
            hapr=hapr,
            pAPR=papr,
            rAPR=rapr,
            mask=np.ones(len(sequence), dtype=np.bool_),
        )
        manifest_records.append({"sequence_id": sid, "length": len(sequence), "file": filename})

    manifest = {
        "format": "unique_sequence_frozen_hapr_v1",
        "checkpoint": str(args.checkpoint),
        "checkpoint_sha256": hashlib.sha256(args.checkpoint.read_bytes()).hexdigest(),
        "boundary_difference_policy": boundary_v3.BOUNDARY_DIFFERENCE_POLICY,
        "source_workbook": str(args.workbook),
        "input_contract": input_contract,
        "sequence_count": len(manifest_records),
        "representation": {
            "hapr": "[L,192] residue-ordered Boundary-Aware Context V3 embedding",
            "pAPR": "[L] APR probability readout",
            "rAPR": "[L] reliability readout",
            "sample_assembly": "deferred; no PDB-chain concatenation in this export",
        },
        "records": manifest_records,
        "model_payload": {
            "variant": payload.get("variant"),
            "refit_epoch": payload.get("refit_epoch"),
            "model_config": payload.get("model_config"),
        },
    }
    (args.output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"event": "complete", "sequence_count": len(manifest_records),
                      "output_dir": str(args.output_dir)}), flush=True)


if __name__ == "__main__":
    main()
